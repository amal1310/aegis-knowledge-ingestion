from __future__ import annotations
import json, re
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
import fitz
from PIL import Image
import pytesseract


def _chunk(text: str, max_chars: int = 1800):
    text = re.sub(r'\s+', ' ', text).strip()
    if not text:
        return []
    out=[]
    for i in range(0, len(text), max_chars):
        out.append(text[i:i+max_chars])
    return out


def parse_pdf(path: Path) -> list[dict[str, Any]]:
    reader = PdfReader(str(path))
    chunks=[]
    doc=fitz.open(str(path))
    for idx, page in enumerate(reader.pages, start=1):
        text=page.extract_text() or ''
        # Scanned pages often have almost no text. OCR the rendered page in that case.
        if len(re.sub(r'\s+', '', text)) < 80:
            pix=doc[idx-1].get_pixmap(matrix=fitz.Matrix(1.6,1.6), alpha=False)
            img=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
            text=pytesseract.image_to_string(img)
            evidence_type='ocr_text'
        else:
            evidence_type='native_text'
        for n,c in enumerate(_chunk(text), start=1):
            chunks.append({'location':f'page {idx}, chunk {n}', 'text':c, 'evidence_type':evidence_type})
    doc.close()
    return chunks


def parse_html(path: Path):
    soup=BeautifulSoup(path.read_text(encoding='utf-8'), 'html.parser')
    chunks=[]
    for i,c in enumerate(_chunk(soup.get_text('\n')),1):
        chunks.append({'location':f'html chunk {i}', 'text':c, 'evidence_type':'native_text'})
    return chunks


def parse_docx(path: Path):
    d=Document(str(path)); chunks=[]
    for i,p in enumerate(d.paragraphs,1):
        if p.text.strip(): chunks.append({'location':f'paragraph {i}', 'text':p.text.strip(), 'evidence_type':'native_text'})
    for ti,t in enumerate(d.tables,1):
        rows=[' | '.join(c.text.strip() for c in r.cells) for r in t.rows]
        chunks.append({'location':f'table {ti}', 'text':'\n'.join(rows), 'evidence_type':'table'})
    return chunks


def parse_xlsx(path: Path):
    wb=load_workbook(str(path), data_only=True)
    chunks=[]
    for ws in wb.worksheets:
        rows=[]
        for row in ws.iter_rows(values_only=True):
            vals=[('' if v is None else str(v)) for v in row]
            if any(vals): rows.append(' | '.join(vals))
        for i,c in enumerate(_chunk('\n'.join(rows)),1):
            chunks.append({'location':f'sheet {ws.title}, chunk {i}', 'text':c, 'evidence_type':'table'})
    return chunks


def parse_json(path: Path):
    data=json.loads(path.read_text(encoding='utf-8'))
    text=json.dumps(data, indent=2)
    return [{'location':'JSON document', 'text':c, 'evidence_type':'structured_json'} for c in _chunk(text)]


def parse_pptx(path: Path):
    prs=Presentation(str(path)); chunks=[]
    for i,slide in enumerate(prs.slides,1):
        texts=[]
        for sh in slide.shapes:
            if hasattr(sh,'text') and sh.text.strip(): texts.append(sh.text.strip())
        if texts: chunks.append({'location':f'slide {i}', 'text':'\n'.join(texts), 'evidence_type':'slide_text'})
    return chunks


def parse_png(path: Path):
    img=Image.open(path)
    text=pytesseract.image_to_string(img)
    return [{'location':'image', 'text':text, 'evidence_type':'ocr_text'}]


def parse_file(path: Path):
    ext=path.suffix.lower()
    return {
        '.pdf':parse_pdf, '.html':parse_html, '.docx':parse_docx, '.xlsx':parse_xlsx,
        '.json':parse_json, '.pptx':parse_pptx, '.png':parse_png,
    }.get(ext, lambda p: [])(path)
